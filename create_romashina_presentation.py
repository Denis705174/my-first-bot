from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("/workspace/Ромашина О.В. презентация.pptx")
SRC_PPTX = Path("/home/ubuntu/.cursor/projects/workspace/uploads/__________._._____________2f4c.pptx")
ASSETS_DIR = Path("/workspace/assets_romashina")
ASSETS_DIR.mkdir(exist_ok=True)


def extract_uploaded_images():
    if SRC_PPTX.exists():
        src_prs = Presentation(SRC_PPTX)
        mapping = {
            (1, 0): "portrait_civil_defense.jpg",
            (2, 0): "desk_front.png",
            (4, 0): "client_consultation.png",
            (5, 0): "desk_side.png",
        }
        for s_idx, slide in enumerate(src_prs.slides, 1):
            p_idx = 0
            for sh in slide.shapes:
                if sh.shape_type == 13:
                    key = (s_idx, p_idx)
                    if key in mapping:
                        (ASSETS_DIR / mapping[key]).write_bytes(sh.image.blob)
                    p_idx += 1


extract_uploaded_images()

PORTRAIT_PHOTO = ASSETS_DIR / "portrait_civil_defense.jpg"
DESK_FRONT_PHOTO = ASSETS_DIR / "desk_front.png"
CLIENT_PHOTO = ASSETS_DIR / "client_consultation.png"
DESK_SIDE_PHOTO = ASSETS_DIR / "desk_side.png"

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

NAVY = RGBColor(20, 42, 114)
BLUE = RGBColor(28, 92, 166)
CYAN = RGBColor(0, 166, 227)
MAGENTA = RGBColor(215, 24, 106)
INK = RGBColor(30, 41, 59)
MUTED = RGBColor(93, 108, 130)
LIGHT = RGBColor(244, 247, 251)
PALE_BLUE = RGBColor(233, 242, 251)
PALE_PINK = RGBColor(252, 234, 242)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(217, 225, 236)


def rect(slide, x, y, w, h, fill, radius=False, line=None, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = fill if line is None else line
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def textbox(
    slide, text, x, y, w, h, size=14, color=INK, bold=False,
    font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
    margin=0, italic=False, line_spacing=1.05,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def rich_text(slide, lines, x, y, w, h, margin=0, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item.get("text", "")
        p.font.name = item.get("font", "Aptos")
        p.font.size = Pt(item.get("size", 14))
        p.font.bold = item.get("bold", False)
        p.font.italic = item.get("italic", False)
        p.font.color.rgb = item.get("color", INK)
        p.alignment = item.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(item.get("after", 4))
        p.space_before = Pt(item.get("before", 0))
        p.line_spacing = item.get("spacing", 1.05)
    return box


def crop_picture(slide, path, x, y, w, h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    with Image.open(path) as im:
        image_ratio = im.width / im.height
    box_ratio = w / h
    if image_ratio > box_ratio:
        visible = box_ratio / image_ratio
        crop = (1 - visible) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        visible = image_ratio / box_ratio
        crop = (1 - visible) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def brand_bar(slide):
    rect(slide, 0, 7.36, 8.7, 0.14, NAVY)
    rect(slide, 8.7, 7.36, 1.6, 0.14, MAGENTA)
    rect(slide, 10.3, 7.36, 3.033, 0.14, CYAN)


# ==============================================================================
# SLIDE 1: Титульный блок + Профессиональный путь + Моя роль и обязанности
# ==============================================================================
s1 = prs.slides.add_slide(blank)
s1.background.fill.solid()
s1.background.fill.fore_color.rgb = LIGHT

# Left panel (Navy)
rect(s1, 0, 0, 4.30, 7.36, NAVY)
rect(s1, 4.18, 0, 0.12, 7.36, MAGENTA)

textbox(s1, "СОЦИАЛЬНЫЙ ФОНД РОССИИ", 0.36, 0.30, 3.60, 0.24, 10, CYAN, True)
textbox(s1, "ОДИН ДЕНЬ РАБОТЫ В СФР", 0.36, 0.56, 3.65, 0.42, 18.5, WHITE, True)
rect(s1, 0.36, 1.04, 0.95, 0.05, MAGENTA)

# Portrait photo in left panel
crop_picture(s1, PORTRAIT_PHOTO, 0.36, 1.20, 3.55, 2.55)
pic_border = rect(s1, 0.36, 1.20, 3.55, 2.55, WHITE, radius=False, line=CYAN, transparency=100)
pic_border.line.width = Pt(1.2)

rich_text(s1, [
    {"text": "Ромашина Ольга Викторовна", "size": 15, "bold": True, "color": WHITE, "after": 3},
    {"text": "Ведущий специалист-эксперт", "size": 11.5, "bold": True, "color": CYAN, "after": 2},
    {"text": "отдела установления мер социальной поддержки семьям с детьми Управления социального обеспечения семей с детьми", "size": 10, "color": RGBColor(220, 229, 245), "spacing": 1.05},
], 0.36, 3.84, 3.60, 1.22)

# Stats badges in left panel
stat1 = rect(s1, 0.36, 5.14, 1.72, 0.90, BLUE, radius=True)
textbox(s1, "19 лет", 0.42, 5.20, 1.60, 0.34, 18, WHITE, True, align=PP_ALIGN.CENTER)
textbox(s1, "профессионального\nопыта", 0.42, 5.54, 1.60, 0.42, 9.5, RGBColor(225, 235, 250), align=PP_ALIGN.CENTER)

stat2 = rect(s1, 2.19, 5.14, 1.72, 0.90, BLUE, radius=True)
textbox(s1, "3,5 года", 2.25, 5.20, 1.60, 0.34, 18, CYAN, True, align=PP_ALIGN.CENTER)
textbox(s1, "в Управлении соц.\nобеспечения семей", 2.25, 5.54, 1.60, 0.42, 9.5, WHITE, align=PP_ALIGN.CENTER)

# Education box in left panel
edu_box = rect(s1, 0.36, 6.14, 3.55, 0.78, RGBColor(14, 31, 88), radius=True, line=BLUE)
edu_box.line.width = Pt(1)
textbox(s1, "ОБРАЗОВАНИЕ", 0.48, 6.20, 3.30, 0.20, 8.5, CYAN, True)
textbox(
    s1,
    "Московский государственный университет экономики, статистики и информатики",
    0.48, 6.40, 3.30, 0.46, 9.5, WHITE, line_spacing=1.05,
)

textbox(s1, "ЗАБОТА • ПОЛЕЗНОСТЬ • НАДЁЖНОСТЬ", 0.36, 7.02, 3.55, 0.22, 8.5, CYAN, True)

# Right side of Slide 1: Professional Path & Focus (Top)
textbox(s1, "ПРОФЕССИОНАЛЬНЫЙ ПУТЬ И ОБЯЗАННОСТИ", 4.58, 0.28, 7.60, 0.24, 10, MAGENTA, True)
textbox(s1, "Опыт, знания и каждое решение — в интересах семьи", 4.58, 0.52, 7.80, 0.42, 20, NAVY, True)
textbox(s1, "01 / 02", 12.10, 0.34, 0.85, 0.28, 10, MUTED, True, align=PP_ALIGN.RIGHT)
rect(s1, 4.58, 1.00, 8.38, 0.02, LINE)

# Focus card + Desk photo
focus_card = rect(s1, 4.58, 1.14, 5.55, 1.48, WHITE, radius=True, line=LINE)
focus_card.line.width = Pt(1)
rect(s1, 4.58, 1.14, 0.09, 1.48, BLUE)
textbox(s1, "ПРОФЕССИОНАЛЬНЫЙ ФОКУС", 4.82, 1.24, 5.10, 0.22, 10, BLUE, True)
textbox(
    s1,
    "Точная оценка права на поддержку, внимание к обстоятельствам каждой семьи и понятная помощь гражданам.",
    4.82, 1.50, 5.12, 0.56, 12, INK, True, line_spacing=1.08,
)
textbox(
    s1,
    "Профессиональный стаж — 19 лет, из них 3,5 года в отделе установления мер социальной поддержки семьям с детьми.",
    4.82, 2.08, 5.12, 0.44, 10.5, MUTED, line_spacing=1.05,
)

crop_picture(s1, DESK_FRONT_PHOTO, 10.30, 1.14, 2.66, 1.48)
photo_border1 = rect(s1, 10.30, 1.14, 2.66, 1.48, WHITE, line=LINE, transparency=100)
photo_border1.line.width = Pt(1)

# Responsibilities header
textbox(s1, "МОЯ РОЛЬ И ОБЯЗАННОСТИ", 4.58, 2.76, 8.38, 0.26, 11, MAGENTA, True)

duties = [
    ("1", "Соблюдение законодательства РФ", "Неукоснительно соблюдать законодательство Российской Федерации при рассмотрении обращений и назначении мер поддержки.", MAGENTA),
    ("2", "Консультирование граждан", "Консультирование граждан о назначении детских пособий, помощь в разъяснении порядка и условий предоставления выплат.", BLUE),
    ("3", "Работа в ГИС ЕЦП и анализ сведений", "Осуществлять в ГИС ЕЦП анализ заявлений, ручную обработку сведений Росреестра, вводить данные из скан-образцов и нецифровых источников подтверждающих документов.", BLUE),
    ("4", "Применение информационных систем", "В своей деятельности использовать программное обеспечение ГИС ЕЦП, АРМ ГС, СЭД, а также разрешенные к применению носители информации.", MAGENTA),
    ("5", "Защита информации и дисциплина", "Обеспечивать выполнение требований по защите информации и соблюдение внутреннего трудового распорядка отделения.", NAVY),
]

for idx, (num, title_txt, desc_txt, accent) in enumerate(duties):
    y = 3.06 + idx * 0.78
    card = rect(s1, 4.58, y, 8.38, 0.68, WHITE, radius=True, line=LINE)
    card.line.width = Pt(1)
    rect(s1, 4.58, y, 0.08, 0.68, accent)
    badge = rect(s1, 4.76, y + 0.14, 0.40, 0.40, accent, radius=True)
    textbox(s1, num, 4.76, y + 0.19, 0.40, 0.30, 11.5, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(s1, title_txt, 5.28, y + 0.08, 7.50, 0.22, 11, NAVY, True)
    textbox(s1, desc_txt, 5.28, y + 0.31, 7.52, 0.34, 9.8, INK, line_spacing=1.02)

textbox(s1, "РОМАШИНА ОЛЬГА ВИКТОРОВНА", 4.58, 7.08, 4.5, 0.20, 8, MUTED, True)
textbox(s1, "1", 12.36, 7.08, 0.6, 0.20, 8, MUTED, True, align=PP_ALIGN.RIGHT)
brand_bar(s1)


# ==============================================================================
# SLIDE 2: Миссия СФР + Ценности Фонда + Рабочий процесс в кадре + Заключение
# ==============================================================================
s2 = prs.slides.add_slide(blank)
s2.background.fill.solid()
s2.background.fill.fore_color.rgb = LIGHT

textbox(s2, "МИССИЯ, ЦЕННОСТИ И ИТОГИ РАБОТЫ", 0.56, 0.28, 8.80, 0.24, 10, MAGENTA, True)
textbox(s2, "Забота о благополучии каждого человека и ориентиры в работе", 0.56, 0.52, 11.20, 0.42, 20, NAVY, True)
textbox(s2, "02 / 02", 12.00, 0.34, 0.78, 0.28, 10, MUTED, True, align=PP_ALIGN.RIGHT)
rect(s2, 0.56, 1.00, 12.21, 0.02, LINE)

# Left column: Mission card (with client photo)
mission_card = rect(s2, 0.56, 1.16, 5.95, 4.16, WHITE, radius=True, line=LINE)
mission_card.line.width = Pt(1)
rect(s2, 0.56, 1.16, 5.95, 0.08, MAGENTA)

crop_picture(s2, CLIENT_PHOTO, 0.76, 1.38, 2.35, 3.74)
textbox(s2, "МИССИЯ СФР", 3.28, 1.36, 3.00, 0.24, 10.5, MAGENTA, True)
textbox(s2, "Мы заботимся\nо благополучии\nкаждого человека", 3.28, 1.62, 3.05, 0.96, 17.5, NAVY, True, line_spacing=1.05)
rect(s2, 3.28, 2.66, 0.85, 0.05, MAGENTA)
textbox(
    s2,
    "Реализуя социальную политику государства, мы обеспечиваем простоту и удобство получения гражданами мер социальной поддержки на основе современных технологий.",
    3.28, 2.82, 3.05, 1.36, 10.8, INK, line_spacing=1.08,
)
mission_note = rect(s2, 3.28, 4.30, 3.05, 0.82, PALE_BLUE, radius=True)
textbox(
    s2,
    "Моя работа помогает воплощать эту миссию каждый день.",
    3.42, 4.44, 2.78, 0.54, 10.5, BLUE, True, line_spacing=1.08,
)

# Right column: Values of the Fund + Desk Side Photo
values_card = rect(s2, 6.72, 1.16, 6.05, 4.16, WHITE, radius=True, line=LINE)
values_card.line.width = Pt(1)
rect(s2, 6.72, 1.16, 6.05, 0.08, BLUE)

textbox(s2, "ЦЕННОСТИ ФОНДА — ОРИЕНТИРЫ В ЕЖЕДНЕВНОЙ РАБОТЕ", 6.94, 1.34, 5.60, 0.24, 10.5, BLUE, True)

values = [
    ("ЗАБОТА", "внимание к человеку", True),
    ("ПОЛЕЗНОСТЬ", "реальная помощь", False),
    ("ДОСТУПНОСТЬ И ОТКРЫТОСТЬ", "понятные решения", False),
    ("НАДЁЖНОСТЬ", "ответственность за результат", True),
    ("САМОСОВЕРШЕНСТВОВАНИЕ", "новые знания и навыки", True),
    ("ЕДИНСТВО", "работа одной командой", False),
]

for i, (v_head, v_body, is_pink) in enumerate(values):
    col, row = i % 2, i // 2
    vx = 6.94 + col * 1.88
    vy = 1.68 + row * 1.14
    fill = PALE_PINK if is_pink else PALE_BLUE
    accent = MAGENTA if is_pink else BLUE
    rect(s2, vx, vy, 1.76, 0.98, fill, radius=True)
    textbox(s2, v_head, vx + 0.08, vy + 0.14, 1.60, 0.36, 8.5, accent, True, align=PP_ALIGN.CENTER, line_spacing=1.0)
    textbox(s2, v_body, vx + 0.08, vy + 0.54, 1.60, 0.32, 9.0, INK, align=PP_ALIGN.CENTER, line_spacing=1.0)

crop_picture(s2, DESK_SIDE_PHOTO, 10.74, 1.68, 1.84, 3.42)
rect(s2, 10.74, 4.38, 1.84, 0.72, NAVY, transparency=12)
textbox(
    s2,
    "Работа в ГИС ЕЦП\nи анализ документов",
    10.80, 4.48, 1.72, 0.52, 8.5, WHITE, True, align=PP_ALIGN.CENTER, line_spacing=1.05,
)

# Bottom Banner: Conclusion & Thank You (Slide 6 content)
footer_banner = rect(s2, 0.56, 5.48, 12.21, 1.50, NAVY, radius=True)
rect(s2, 0.56, 5.48, 0.12, 1.50, MAGENTA)

textbox(s2, "ОДИН ДЕНЬ РАБОТЫ В СФР • ЗАКЛЮЧЕНИЕ", 0.88, 5.62, 7.40, 0.22, 9.5, CYAN, True)
textbox(
    s2,
    "Опыт помогает принимать взвешенные решения, а ценности Фонда — сохранять внимание к каждому человеку.",
    0.88, 5.90, 7.55, 0.56, 13.5, WHITE, False, line_spacing=1.10,
)
textbox(
    s2,
    "Забота о людях — основа нашей работы.",
    0.88, 6.52, 7.55, 0.30, 12, RGBColor(202, 225, 250), True, italic=True,
)

thanks_box = rect(s2, 8.68, 5.68, 3.86, 1.10, MAGENTA, radius=True)
textbox(
    s2,
    "СПАСИБО ЗА ВНИМАНИЕ!",
    8.78, 5.86, 3.66, 0.40, 18, WHITE, True, align=PP_ALIGN.CENTER,
)
textbox(
    s2,
    "Ромашина Ольга Викторовна",
    8.78, 6.32, 3.66, 0.28, 10.5, WHITE, False, align=PP_ALIGN.CENTER,
)

textbox(s2, "РОМАШИНА ОЛЬГА ВИКТОРОВНА", 0.56, 7.08, 4.5, 0.20, 8, MUTED, True)
textbox(s2, "2", 12.17, 7.08, 0.6, 0.20, 8, MUTED, True, align=PP_ALIGN.RIGHT)
brand_bar(s2)

prs.core_properties.title = "Один день работы в СФР — Ромашина Ольга Викторовна (2 слайда)"
prs.core_properties.subject = "Профессиональная презентация сотрудника СФР"
prs.core_properties.author = "Ромашина Ольга Викторовна"
prs.core_properties.keywords = "СФР, социальная поддержка, семьи с детьми"
prs.save(OUT)
print(OUT)
